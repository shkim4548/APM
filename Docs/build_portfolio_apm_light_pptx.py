# -*- coding: utf-8 -*-
"""
APM Agent & Console 포트폴리오 PPT — 라이트 테마(Docs/portfolio_apm_light.html) 기반, python-pptx 네이티브 생성.
Marp+LibreOffice 경로에서 텍스트 런 경계의 숫자 소실 버그가 발견되어(예: "AES-256-GCM"→"AES- -GCM"),
HTML→LibreOffice 변환 단계 자체가 없는 이 스크립트로 대체.
"""
import os
import math
from PIL import Image as PILImage
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ── 색상 팔레트 (Docs/portfolio_apm_light.html :root 그대로) ──
BG        = RGBColor(0xFF, 0xFF, 0xFF)
CARD      = RGBColor(0xF6, 0xF8, 0xFA)
CARD2     = RGBColor(0xEA, 0xEE, 0xF2)
BORDER    = RGBColor(0xD0, 0xD7, 0xDE)
BORDER2   = RGBColor(0xE4, 0xE8, 0xEC)
ACCENT    = RGBColor(0x09, 0x69, 0xDA)
SUCCESS   = RGBColor(0x1A, 0x7F, 0x37)
WARNING   = RGBColor(0x9A, 0x67, 0x00)
DANGER    = RGBColor(0xCF, 0x22, 0x2E)
TEXT1     = RGBColor(0x1F, 0x23, 0x28)
TEXT2     = RGBColor(0x59, 0x63, 0x6E)
TEXT3     = RGBColor(0x6E, 0x77, 0x81)
CALLOUT_BG = RGBColor(0xEA, 0xF6, 0xEC)
CALLOUT_BORDER = RGBColor(0xAF, 0xDC, 0xB8)

# code-block syntax colors — matches Docs/portfolio_apm_light.html .c-key/.c-type/.c-fn/.c-str
CODE_KEY  = DANGER
CODE_TYPE = RGBColor(0x05, 0x50, 0xAE)
CODE_FN   = RGBColor(0x82, 0x50, 0xDF)
CODE_STR  = RGBColor(0x11, 0x63, 0x29)
CODE_CMT  = TEXT3

# chart series colors — dataviz skill reference palette slots 1(blue)/2(orange),
# validated as an adjacent categorical pair (CVD ΔE 24.7, normal-vision ΔE 33.6)
CHART_BLUE   = RGBColor(0x2A, 0x78, 0xD6)
CHART_ORANGE = RGBColor(0xEB, 0x68, 0x34)

FONT_MONO = "Noto Sans Mono CJK KR"
FONT_MONO_LATIN = "Noto Sans Mono"  # non-CJK mono: proper narrow fixed-width for Latin/digits,
                                    # avoids the CJK mono font's full-width-cell stretching on ASCII text
FONT_BODY = "Noto Sans CJK KR"


def _is_wide(ch):
    return ord(ch) > 0x2E7F


def _script_segments(s):
    """Split s into consecutive runs of (text, is_wide) — CJK vs everything else."""
    if not s:
        return [(s, False)]
    segs = []
    cur = s[0]
    cur_wide = _is_wide(s[0])
    for ch in s[1:]:
        w = _is_wide(ch)
        if w == cur_wide:
            cur += ch
        else:
            segs.append((cur, cur_wide))
            cur, cur_wide = ch, w
    segs.append((cur, cur_wide))
    return segs


def _add_text_runs(p, s, size, color, bold, font):
    """Add run(s) for s to paragraph p. For FONT_MONO, splits by script so Latin/digits
    use the narrow non-CJK mono font instead of the CJK mono font's full-width cells."""
    if font == FONT_MONO:
        for seg, wide in _script_segments(s):
            r = p.add_run()
            r.text = seg
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.name = FONT_MONO if wide else FONT_MONO_LATIN
            r.font.color.rgb = color
    else:
        r = p.add_run()
        r.text = s
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.name = font
        r.font.color.rgb = color

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

prs = Presentation()
prs.slide_width = SLIDE_W
prs.slide_height = SLIDE_H
BLANK = prs.slide_layouts[6]


def new_slide():
    s = prs.slides.add_slide(BLANK)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG
    bg.line.fill.background()
    bg.shadow.inherit = False
    s.shapes._spTree.remove(bg._element)
    s.shapes._spTree.insert(2, bg._element)
    return s


def rect(slide, x, y, w, h, fill=None, line_color=None, line_w=0.75, radius=0.06, shadow=False):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    try:
        shp.adjustments[0] = radius
    except Exception:
        pass
    if fill is not None:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    else:
        shp.fill.background()
    if line_color is not None:
        shp.line.color.rgb = line_color
        shp.line.width = Pt(line_w)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = shadow
    return shp


def text(slide, x, y, w, h, s, size, color, bold=False, font=FONT_BODY,
         align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, line_spacing=1.25, wrap=True):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = line_spacing
    _add_text_runs(p, s, size, color, bold, font)
    return box


def multi_run(slide, x, y, w, h, runs, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP,
              line_spacing=1.3, wrap=True):
    """runs: list of (text, color, bold, font, size)"""
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = line_spacing
    for (t, color, bold, font, size) in runs:
        _add_text_runs(p, t, size, color, bold, font)
    return box


def _text_width_in(s, size=9.5):
    # CJK glyphs render roughly full-width (~2x) vs ASCII in a mono CJK font
    wide = sum(1 for ch in s if ord(ch) > 0x2E7F)
    narrow = len(s) - wide
    per_narrow = size * 0.0092
    per_wide = size * 0.018
    return narrow * per_narrow + wide * per_wide


def _est_lines(s, size, avail_w_in):
    """Wrap-line estimate using the CJK/Latin-aware width function (not a flat
    chars-per-line count, which underestimates lines for Korean-heavy strings
    and previously caused mini-card text to overflow/clip)."""
    if avail_w_in <= 0:
        return 1
    return max(1, math.ceil(_text_width_in(s, size) / (avail_w_in * 0.95)))


def lines_block(slide, x, y, w, h, lines, font=FONT_MONO, size=11, color=TEXT2,
                 line_spacing=1.25, align=PP_ALIGN.LEFT):
    """lines: list of (text, bold) or plain str; each becomes its own paragraph."""
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    for i, item in enumerate(lines):
        t, bold = item if isinstance(item, tuple) else (item, False)
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        _add_text_runs(p, t, size, color, bold, font)
    return box


def tag(slide, x, y, s, color, border_color, size=10.5):
    w = Inches(0.28 + _text_width_in(s, size))
    h = Inches(0.28)
    box = rect(slide, x, y, w, h, fill=CARD, line_color=border_color, line_w=0.75, radius=0.35)
    text(slide, x, y, w, h, s, size, color, font=FONT_MONO,
         align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, wrap=False)
    return w


def page_num(slide, n, total=18):
    text(slide, Inches(12.55), Inches(7.12), Inches(0.65), Inches(0.28),
         f"{n} / {total}", 10.5, TEXT3, font=FONT_MONO, align=PP_ALIGN.RIGHT)


def title_block(slide, kicker, title, sub=None):
    y = Inches(0.35)
    if kicker:
        text(slide, Inches(0.55), y, Inches(10), Inches(0.28), kicker, 11, TEXT3, font=FONT_MONO)
        y = y + Inches(0.32)
    text(slide, Inches(0.55), y, Inches(12.2), Inches(0.55), title, 24, TEXT1, bold=True, font=FONT_MONO)
    y = y + Inches(0.55)
    if sub:
        multi_run(slide, Inches(0.55), y, Inches(12.2), Inches(0.5), sub, line_spacing=1.2)
        y = y + Inches(0.42)
    return y


def sub_line(runs_or_text, size=13):
    """helper to build a single sub-text run list from plain string or list of (t,color,bold)"""
    if isinstance(runs_or_text, str):
        return [(runs_or_text, TEXT2, False, FONT_BODY, size)]
    out = []
    for (t, color, bold) in runs_or_text:
        out.append((t, color, bold, FONT_MONO if bold else FONT_BODY, size))
    return out


DECISION_LABEL_COLORS = {"문제": TEXT3, "판단": ACCENT, "해결": SUCCESS, "원인": ACCENT, "재실측": ACCENT}


def decision_card(slide, x, y, w, num, title_s, rows):
    """Design-decision card for the '설계 결정' slides: 문제/판단/해결 rows split into their
    own bordered mini-cards (instead of stacked label+text lines) for readability."""
    w_in = w / Inches(1)
    inner_w_in = w_in - 0.48
    text_w_in = inner_w_in - 0.32  # mini-card's own left+right padding

    top_h_in = 0.16 + (0.28 if num else 0) + 0.40
    mini_heights_in = []
    for (_, body, _color) in rows:
        lines = _est_lines(body, 11, text_w_in)
        mini_heights_in.append(0.34 + 0.19 * lines)
    total_h_in = top_h_in + sum(mini_heights_in) + 0.12 * len(rows) + 0.14
    h = Inches(total_h_in)

    rect(slide, x, y, w, h, fill=CARD, line_color=BORDER, line_w=0.75, radius=0.045)
    cy = y + Inches(0.16)
    if num:
        text(slide, x + Inches(0.24), cy, w - Inches(0.48), Inches(0.22), num, 10.5, TEXT3, font=FONT_MONO)
        cy += Inches(0.28)
    text(slide, x + Inches(0.24), cy, w - Inches(0.48), Inches(0.3), title_s, 14.5, TEXT1, bold=True, font=FONT_MONO)
    cy += Inches(0.40)

    inner_x = x + Inches(0.24)
    inner_w = w - Inches(0.48)
    for (label, body, color), mini_h_in in zip(rows, mini_heights_in):
        mini_h = Inches(mini_h_in)
        label_color = DECISION_LABEL_COLORS.get(label, TEXT3)
        body_color = color if color else TEXT2
        rect(slide, inner_x, cy, inner_w, mini_h, fill=BG, line_color=BORDER2, line_w=0.75, radius=0.09)
        text(slide, inner_x + Inches(0.16), cy + Inches(0.09), inner_w - Inches(0.32), Inches(0.2),
             label, 10, label_color, bold=True, font=FONT_MONO)
        multi_run(slide, inner_x + Inches(0.16), cy + Inches(0.30), inner_w - Inches(0.32), mini_h - Inches(0.36),
                  [(body, body_color, False, FONT_BODY, 11)], line_spacing=1.25)
        cy += mini_h + Inches(0.12)
    return h


def bug_fix(slide, x, y, w, kind, header, chunks):
    """HTML technical-review 'bug-card' style: plain card bg, colored accent limited to
    a left bar + header text. Body is split into labeled mini-cards (chunks: list of
    (label, text, color_or_None)) instead of one long run-on paragraph."""
    is_bug = kind == "bug"
    accent = DANGER if is_bug else SUCCESS
    mark = "✕ " if is_bug else "✓ "

    mini_w_in = w / Inches(1) - 0.5
    text_w_in = mini_w_in - 0.32  # mini-card's own left+right padding
    mini_heights_in = []
    for (_, chunk_text, _color) in chunks:
        lines = _est_lines(chunk_text, 11, text_w_in)
        mini_heights_in.append(0.34 + 0.19 * lines)
    header_h_in = 0.5
    h = Inches(header_h_in + sum(mini_heights_in) + 0.1 * len(chunks) + 0.14)

    rect(slide, x, y, w, h, fill=CARD, line_color=BORDER, line_w=0.75, radius=0.04)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y + Inches(0.05), Inches(0.06), h - Inches(0.1))
    bar.fill.solid()
    bar.fill.fore_color.rgb = accent
    bar.line.fill.background()
    bar.shadow.inherit = False
    multi_run(slide, x + Inches(0.26), y + Inches(0.14), w - Inches(0.5), Inches(0.3),
              [(mark + header, accent, True, FONT_MONO, 13)])

    cy = y + Inches(header_h_in)
    mini_x, mini_w = x + Inches(0.26), w - Inches(0.5)
    for (label, chunk_text, color), mini_h_in in zip(chunks, mini_heights_in):
        mini_h = Inches(mini_h_in)
        rect(slide, mini_x, cy, mini_w, mini_h, fill=BG, line_color=BORDER2, line_w=0.75, radius=0.09)
        text(slide, mini_x + Inches(0.16), cy + Inches(0.09), mini_w - Inches(0.32), Inches(0.2),
             label, 10, TEXT3, bold=True, font=FONT_MONO)
        multi_run(slide, mini_x + Inches(0.16), cy + Inches(0.30), mini_w - Inches(0.32), mini_h - Inches(0.36),
                  [(chunk_text, color if color else TEXT2, False, FONT_BODY, 11)], line_spacing=1.25)
        cy += mini_h + Inches(0.1)
    return h


def arch_box(slide, x, y, w, h, accent, title, subtitle, bullets):
    """Pipeline stage box for the architecture diagram — same card+left-accent-bar
    language as bug_fix(), so the diagram reads as part of the same design system
    instead of a one-off graphic."""
    rect(slide, x, y, w, h, fill=CARD, line_color=BORDER, line_w=0.75, radius=0.05)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y + Inches(0.05), Inches(0.07), h - Inches(0.1))
    bar.fill.solid()
    bar.fill.fore_color.rgb = accent
    bar.line.fill.background()
    bar.shadow.inherit = False
    ty = y + Inches(0.18)
    text(slide, x + Inches(0.26), ty, w - Inches(0.5), Inches(0.32), title, 16, accent, bold=True, font=FONT_MONO)
    ty += Inches(0.36)
    if subtitle:
        text(slide, x + Inches(0.26), ty, w - Inches(0.5), Inches(0.24), subtitle, 10.5, TEXT3, font=FONT_MONO)
        ty += Inches(0.3)
    ty += Inches(0.08)
    lines = [("· " + b, False) for b in bullets]
    lines_block(slide, x + Inches(0.26), ty, w - Inches(0.5), h - (ty - y) - Inches(0.15), lines,
                font=FONT_BODY, size=11, color=TEXT2, line_spacing=1.32)


def arch_arrow(slide, x, box_y, w, box_h, label, label_y, label_h):
    """Connector between two arch_box() stages. The caption lives in a dedicated band
    above the boxes (label_y/label_h) so it never overlaps a box title or its bullets."""
    text(slide, x - Inches(0.55), label_y, w + Inches(1.1), label_h, label, 9.5, TEXT3,
         font=FONT_MONO, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.BOTTOM, wrap=False)
    ay = box_y + Inches(0.52)
    ah = Inches(0.26)
    arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x, ay, w, ah)
    arrow.fill.solid()
    arrow.fill.fore_color.rgb = BORDER
    arrow.line.fill.background()
    arrow.shadow.inherit = False
    try:
        arrow.adjustments[0] = 0.55
        arrow.adjustments[1] = 0.5
    except Exception:
        pass


SCREENSHOT_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "screenshots")


def screenshot_card(slide, x, y, w, h, label, filename):
    """Label + real screenshot, letterboxed (not cropped) inside a bordered mat so
    dashboard/chart content never gets cut off by an aspect-ratio mismatch."""
    text(slide, x, y, w, Inches(0.28), label, 13, TEXT1, bold=True, font=FONT_MONO)
    mat_y = y + Inches(0.34)
    mat_h = h - Inches(0.34)
    rect(slide, x, mat_y, w, mat_h, fill=CARD, line_color=BORDER, line_w=0.75, radius=0.04)
    path = os.path.join(SCREENSHOT_DIR, filename)
    iw, ih = PILImage.open(path).size
    pad = Inches(0.08)
    avail_w, avail_h = w - pad * 2, mat_h - pad * 2
    scale = min(avail_w / iw, avail_h / ih)
    disp_w, disp_h = int(iw * scale), int(ih * scale)
    img_x = x + (w - disp_w) // 2
    img_y = mat_y + (mat_h - disp_h) // 2
    slide.shapes.add_picture(path, img_x, img_y, width=disp_w, height=disp_h)


def tech_card(slide, x, y, w, title_s, body_str):
    """Tech-stack card: title + its '·'-separated items as flow-wrapped mini-cards
    underneath, instead of one run-on line of text."""
    items = [it.strip() for it in body_str.split("·") if it.strip()]
    pad_in = 0.22
    avail_w_in = w / Inches(1) - 2 * pad_in
    mini_h_in = 0.34
    gap_in = 0.1
    item_size = 10.5

    rows, cur_row, cur_w = [], [], 0.0
    for it in items:
        iw_in = 0.26 + _text_width_in(it, item_size)
        if cur_row and cur_w + gap_in + iw_in > avail_w_in:
            rows.append(cur_row)
            cur_row, cur_w = [], 0.0
        cur_row.append((it, iw_in))
        cur_w += iw_in + gap_in
    if cur_row:
        rows.append(cur_row)

    title_h_in = 0.32
    content_h_in = pad_in + title_h_in + 0.06 + len(rows) * mini_h_in + max(0, len(rows) - 1) * gap_in + pad_in
    h = Inches(content_h_in)

    rect(slide, x, y, w, h, fill=CARD, line_color=BORDER, line_w=0.75, radius=0.05)
    text(slide, x + Inches(pad_in), y + Inches(pad_in), w - Inches(2 * pad_in), Inches(title_h_in),
         title_s, 13, TEXT1, bold=True, font=FONT_MONO)
    ry = y + Inches(pad_in + title_h_in + 0.06)
    for row in rows:
        rx = x + Inches(pad_in)
        for (it, iw_in) in row:
            iw = Inches(iw_in)
            rect(slide, rx, ry, iw, Inches(mini_h_in), fill=BG, line_color=BORDER2, line_w=0.75, radius=0.28)
            text(slide, rx, ry, iw, Inches(mini_h_in), it, item_size, TEXT2, font=FONT_MONO,
                 align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, wrap=False)
            rx += iw + Inches(gap_in)
        ry += Inches(mini_h_in + gap_in)
    return h


def code_block(slide, x, y, w, lines, size=11, line_spacing=1.3):
    """lines: list of paragraphs, each a list of (text, color_or_None) token tuples.
    Real syntax-tinted code excerpt (matches portfolio_apm_light.html .code-block)."""
    pad_in = 0.16
    line_h_in = size / 72 * line_spacing * 1.35
    h = Inches(pad_in * 2 + line_h_in * len(lines))
    rect(slide, x, y, w, h, fill=CARD2, line_color=BORDER, line_w=0.75, radius=0.035)
    box = slide.shapes.add_textbox(x + Inches(pad_in), y + Inches(pad_in - 0.04),
                                    w - Inches(pad_in * 2), h - Inches(pad_in * 2 - 0.08))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.line_spacing = line_spacing
        if not line:
            _add_text_runs(p, " ", size, TEXT1, False, FONT_MONO)
            continue
        for (t, color) in line:
            _add_text_runs(p, t, size, color if color else TEXT1, False, FONT_MONO)
    return h


def bar_chart_two_series(slide, x, y, w, caption, categories, s1, s2, label1, label2,
                          color1, color2, max_val=100):
    """Grouped horizontal bar chart — 2 series per category (pill-ended bars, direct
    value labels, swatch legend). Values are 0..max_val on a single shared axis."""
    x_in = x / Inches(1)
    w_in = w / Inches(1)
    label_w_in = 1.5
    val_w_in = 0.5
    bar_area_in = w_in - label_w_in - val_w_in
    row_h_in = 0.42
    bar_h_in = 0.13
    bar_gap_in = 0.03

    text(slide, x, y, Inches(6), Inches(0.22), caption, 10.5, TEXT3, font=FONT_MONO)
    ly_in = (y / Inches(1)) + 0.30
    sw = Inches(0.12)
    rect(slide, Inches(x_in + label_w_in), Inches(ly_in), sw, sw, fill=color1)
    text(slide, Inches(x_in + label_w_in + 0.2), Inches(ly_in - 0.03), Inches(1.3), Inches(0.2),
         label1, 10, TEXT2, font=FONT_MONO)
    rect(slide, Inches(x_in + label_w_in + 1.7), Inches(ly_in), sw, sw, fill=color2)
    text(slide, Inches(x_in + label_w_in + 1.9), Inches(ly_in - 0.03), Inches(1.3), Inches(0.2),
         label2, 10, TEXT2, font=FONT_MONO)

    ry_in = ly_in + 0.32
    for cat, v1, v2 in zip(categories, s1, s2):
        text(slide, x, Inches(ry_in), Inches(label_w_in - 0.1), Inches(row_h_in), cat, 10.5, TEXT2,
             font=FONT_MONO, valign=MSO_ANCHOR.MIDDLE)
        bar_x_in = x_in + label_w_in
        b1w_in = max(0.03, bar_area_in * v1 / max_val)
        b2w_in = max(0.03, bar_area_in * v2 / max_val)
        pair_h_in = bar_h_in * 2 + bar_gap_in
        b1_y_in = ry_in + (row_h_in - pair_h_in) / 2
        b2_y_in = b1_y_in + bar_h_in + bar_gap_in
        rect(slide, Inches(bar_x_in), Inches(b1_y_in), Inches(b1w_in), Inches(bar_h_in),
             fill=color1, radius=0.5)
        rect(slide, Inches(bar_x_in), Inches(b2_y_in), Inches(b2w_in), Inches(bar_h_in),
             fill=color2, radius=0.5)
        text(slide, Inches(bar_x_in + b1w_in + 0.07), Inches(b1_y_in - 0.03), Inches(val_w_in), Inches(bar_h_in + 0.06),
             f"{v1:.0f}%", 9, TEXT1, font=FONT_MONO, valign=MSO_ANCHOR.MIDDLE)
        text(slide, Inches(bar_x_in + b2w_in + 0.07), Inches(b2_y_in - 0.03), Inches(val_w_in), Inches(bar_h_in + 0.06),
             f"{v2:.0f}%", 9, TEXT1, font=FONT_MONO, valign=MSO_ANCHOR.MIDDLE)
        ry_in += row_h_in
    return Inches(ry_in) - y


def stat_block(slide, x, y, val, label, w=Inches(3.6)):
    text(slide, x, y, w, Inches(0.42), val, 26, ACCENT, bold=True, font=FONT_MONO)
    text(slide, x, y + Inches(0.52), w, Inches(0.45), label, 11, TEXT3, line_spacing=1.15)


def callout(slide, x, y, w, h, runs):
    rect(slide, x, y, w, h, fill=CALLOUT_BG, line_color=CALLOUT_BORDER, line_w=0.75, radius=0.06)
    multi_run(slide, x + Inches(0.24), y, w - Inches(0.48), h,
              runs, valign=MSO_ANCHOR.MIDDLE, line_spacing=1.3)


def table(slide, x, y, w, headers, rows, col_widths, row_h=Inches(0.34), header_h=Inches(0.3),
          header_size=10.5, cell_size=11.5):
    cy = y
    cx = x
    for hd, cw in zip(headers, col_widths):
        text(slide, cx, cy, cw, header_h, hd, header_size, TEXT3, font=FONT_MONO)
        cx += cw
    cy += header_h + Inches(0.06)
    line = slide.shapes.add_connector(1, x, cy, x + w, cy)
    line.line.color.rgb = BORDER
    line.line.width = Pt(0.75)
    cy += Inches(0.08)
    for row in rows:
        cx = x
        for cell, cw in zip(row, col_widths):
            if isinstance(cell, tuple):
                s, color = cell
            else:
                s, color = cell, TEXT2
            text(slide, cx, cy, cw, row_h, s, cell_size, color,
                 font=FONT_MONO if color in (SUCCESS, DANGER) else FONT_BODY)
            cx += cw
        cy += row_h
        line2 = slide.shapes.add_connector(1, x, cy - Inches(0.02), x + w, cy - Inches(0.02))
        line2.line.color.rgb = BORDER2
        line2.line.width = Pt(0.5)
    return cy


# ══════════════════════════════════════════════════════════════════
# Slide 1 — Title
# ══════════════════════════════════════════════════════════════════
s = new_slide()
cy = Inches(0.85)
badge_w = Inches(0.12 + len("전체 파이프라인 end-to-end 동작 · 실측 검증 완료") * 0.11)
rect(s, Inches(0.55), cy, badge_w, Inches(0.34), fill=RGBColor(0xE8, 0xF5, 0xEC), line_color=RGBColor(0xB7, 0xDE, 0xC2), line_w=0.75, radius=0.4)
text(s, Inches(0.55), cy, badge_w, Inches(0.34), "전체 파이프라인 end-to-end 동작 · 실측 검증 완료", 11, SUCCESS,
     font=FONT_MONO, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, wrap=False)
cy += Inches(0.62)
text(s, Inches(0.55), cy, Inches(11.5), Inches(0.65), "APM Agent & Console", 36, ACCENT, bold=True, font=FONT_MONO)
cy += Inches(0.66)
text(s, Inches(0.55), cy, Inches(11.5), Inches(0.6), "경량 모니터링 에이전트 + 실시간 웹 대시보드", 27, TEXT1, bold=True, font=FONT_MONO)
cy += Inches(0.68)
multi_run(s, Inches(0.55), cy, Inches(11.2), Inches(0.95), sub_line(
    "기존 IOCP 기반 TCP 서버 경험을 Standalone Asio로 재구성한 크로스플랫폼 네트워크 코어 위에, "
    "TLS·페이로드 이중 암호화(AES-256-GCM)·메시지 프레이밍·타입 안전 직렬화를 갖춘 수집 에이전트와, "
    ".NET 8 플러그인 아키텍처 기반 실시간 대시보드까지 1인 설계·구현.", size=14.5), line_spacing=1.4)
cy += Inches(1.05)

# flow-wrapped tags — single unified set instead of fixed-size rows, so the
# last row never ends up with one lonely tag (previous layout's issue).
hero_tags = [
    ("C++20 / Standalone Asio", ACCENT),
    ("TLS + AES-256-GCM (AEAD)", ACCENT),
    ("Protobuf", TEXT2),
    ("SQLite / TimescaleDB (선택형)", TEXT2),
    (".NET 8 / ASP.NET Core", SUCCESS),
    ("SignalR + Chart.js", SUCCESS),
    ("크로스플랫폼: Linux + Windows", ACCENT),
    ("부하 테스트 + strace 프로파일링", TEXT2),
]
tx, ty = Inches(0.55), cy
right_edge = Inches(12.75)
for label, color in hero_tags:
    tw_est = Inches(0.28 + _text_width_in(label, 10.5))
    if tx > Inches(0.55) and tx + tw_est > right_edge:
        tx, ty = Inches(0.55), ty + Inches(0.44)
    tw = tag(s, tx, ty, label, color, BORDER if color == TEXT2 else color)
    tx += tw + Inches(0.14)
cy = ty + Inches(0.72)

divider = s.shapes.add_connector(1, Inches(0.55), cy, Inches(12.75), cy)
divider.line.color.rgb = BORDER2
divider.line.width = Pt(0.75)
cy += Inches(0.32)
hero_stats = [("5", "수집 지표"), ("300", "동시 접속 부하 테스트"), ("27", "테스트 통과"), ("10", "버그 발견/수정")]
for i, (val, label) in enumerate(hero_stats):
    stat_block(s, Inches(0.55) + Inches(2.95) * i, cy, val, label, w=Inches(2.7))
page_num(s, 1)

# ══════════════════════════════════════════════════════════════════
# Slide 2 — 프로젝트 한눈에 보기
# ══════════════════════════════════════════════════════════════════
s = new_slide()
cy = title_block(s, None, "프로젝트 한눈에 보기")
cy += Inches(0.15)
stats = [
    ("5", "수집 지표 종류(CPU/메모리/디스크/네트워크/TCP)"),
    ("27", "단위 테스트 전부 통과(C++ 9 + .NET 18)"),
    ("300", "동시 접속 부하 테스트(자체 제작 LoadTester)"),
    ("13.8MB", "RSS 5분간 완전 고정(누수 없음)"),
    ("2 / 2", "OS 빌드+실행 검증(Linux, Windows)"),
    ("10", "실행 검증 중 발견·수정한 버그"),
]
card_w, card_h = Inches(3.9), Inches(1.4)
gap_x, gap_y = Inches(0.2), Inches(0.2)
row_y = [cy, cy + card_h + gap_y]
for i, (val, label) in enumerate(stats):
    col = i % 3
    row = i // 3
    fx = Inches(0.55) + (card_w + gap_x) * col
    fy = row_y[row]
    rect(s, fx, fy, card_w, card_h, fill=CARD, line_color=BORDER, line_w=0.75, radius=0.05)
    stat_block(s, fx + Inches(0.24), fy + Inches(0.22), val, label, w=card_w - Inches(0.48))
callout(s, Inches(0.55), Inches(6.95 - 0.55), Inches(12.2), Inches(0.55), sub_line([
    ('"가볍다"는 주장을 수치로 뒷받침 — 모든 지표는 ', TEXT2, False),
    ("strace/proc 실측", TEXT1, True),
    ("과 ", TEXT2, False),
    ("실제 실행 검증", TEXT1, True),
    (" 결과, 추정치가 아니다.", TEXT2, False),
], size=13))
page_num(s, 2)

# ══════════════════════════════════════════════════════════════════
# Slide 3 — 전체 시스템 구조
# ══════════════════════════════════════════════════════════════════
s = new_slide()
cy = title_block(s, "Architecture", "전체 시스템 구조", sub_line(
    "PROJECT_TECHNICAL_REVIEW.md §1 전체 파이프라인 — Agent → Collector → APM_Console, 3개 프로세스."))
cy += Inches(0.15)
label_h = Inches(0.28)
box_y = cy + label_h + Inches(0.06)
diagram_h = Inches(3.6)
box_w = Inches(3.65)
arrow_w = Inches(0.65)
bx = Inches(0.55)
arch_box(s, bx, box_y, box_w, diagram_h, ACCENT, "Agent", None, [
    "ResourceCollector — CPU/메모리/디스크/네트워크/TCP 수집",
    "Protobuf 직렬화(apm::Metric)",
    "PacketHeader{size,id} 프레이밍",
    "AesGcmPayload::Seal() — AES-256-GCM",
])
bx += box_w
arch_arrow(s, bx, box_y, arrow_w, diagram_h, "TLS(asio::ssl)", cy, label_h)
bx += arrow_w
arch_box(s, bx, box_y, box_w, diagram_h, DANGER, "Collector", "asio::io_context — 비동기 I/O 코어", [
    "PacketHandler → 타입 안전 디스패치",
    "IMetricStore — SQLite / TimescaleDB",
    "ResilientSender — 재연결 큐잉 재전송",
    "AesGcmPayload::Seal() — AES-256-GCM",
])
bx += box_w
arch_arrow(s, bx, box_y, arrow_w, diagram_h, "TLS(전용 인증서)", cy, label_h)
bx += arrow_w
arch_box(s, bx, box_y, box_w, diagram_h, WARNING, "APM_Console", "ApmConsole.Host — .NET 8", [
    "MetricsReceiverService — TcpListener+SslStream",
    "AesGcm 복호화 → Metric.Parser.ParseFrom",
    "ApmDbContext(SQLite) 저장",
    "MetricsHub(SignalR) 실시간 브로드캐스트",
    "/apm/dashboard — Chart.js 그래프 5개",
])
callout(s, Inches(0.55), Inches(6.95 - 0.85), Inches(12.2), Inches(0.85), sub_line([
    ("핵심 — Collector는 두 역할을 동시에 한다", TEXT1, True),
    (": (1) Agent 데이터를 로컬 SQLite에 저장하는 \"수집 서버\", (2) 저장한 데이터를 APM_Console로 넘기는 \"중계자\". "
     "이 둘은 서로 독립적 — 저장이 실패해도 전송은 별개로 진행되고, 전송(WebServer 연결)이 끊겨도 로컬 저장은 계속된다.",
     TEXT2, False),
], size=12.5))
page_num(s, 3)

# ══════════════════════════════════════════════════════════════════
# Slide 4 — Collector 내부 구조 (딥다이브)
# ══════════════════════════════════════════════════════════════════
s = new_slide()
cy = title_block(s, "Collector Deep-Dive", "Collector 내부 구조 — 스레드 배치", sub_line(
    "5라운드에 걸친 스케일 개선(9~13p)이 실제로 어떤 스레드 배치로 귀결됐는지 — 공유 상태는 정확히 한 스레드만 건드리게 설계."))
cy += Inches(0.12)
top_h = Inches(1.25)
arch_box(s, Inches(0.55), cy, Inches(12.2), top_h, ACCENT, "네트워크 스레드(메인) — asio::io_context", None, [
    "accept · TLS 핸드셰이크 · PacketHandler 디스패치",
    "저장/로깅은 직접 실행하지 않고 Push()만 하고 즉시 리턴",
])
arrow_y = cy + top_h + Inches(0.04)
for ax, lbl in [(Inches(3.15), "Push(저장 요청)"), (Inches(9.05), "Push(로그 요청)")]:
    text(s, ax - Inches(0.75), arrow_y, Inches(1.9), Inches(0.2), lbl, 9.5, TEXT3,
         font=FONT_MONO, align=PP_ALIGN.CENTER, wrap=False)
    arrow = s.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, ax, arrow_y + Inches(0.2), Inches(0.3), Inches(0.24))
    arrow.fill.solid(); arrow.fill.fore_color.rgb = BORDER
    arrow.line.fill.background(); arrow.shadow.inherit = False
worker_y = arrow_y + Inches(0.5)
worker_h = Inches(1.7)
arch_box(s, Inches(0.55), worker_y, Inches(5.95), worker_h, SUCCESS, "metricStoreQueue 워커 스레드", None, [
    "SqliteMetricStore::Store() 실행 — SQLite WAL 모드",
    "Round 2에서 신설(1-7-b) — 네트워크 스레드 블로킹 해소",
])
arch_box(s, Inches(6.8), worker_y, Inches(5.95), worker_h, SUCCESS, "consoleLogQueue 워커 스레드", None, [
    "콘솔 로그 조립+출력 전담",
    "런타임 중 유일하게 stdout/stderr를 건드리는 스레드",
    "Round 4에서 신설(1-7-e/f)",
])
callout(s, Inches(0.55), Inches(6.95 - 0.95), Inches(12.2), Inches(0.95), sub_line([
    ("설계 원칙 — 공유 상태는 정확히 한 스레드만", TEXT1, True),
    (" — stdin은 별도 스레드로 읽되 실제 전송은 asio::post()로 네트워크 스레드에 위임(락 불필요). "
     "단, 워커가 1개→3개로 늘며 스레드 간 락 경합(futex)이 300-agent 규모의 새 병목으로 떠올랐다(9~12p 참고) — "
     "트레이드오프 없는 개선은 없었다.", TEXT2, False),
], size=12))
page_num(s, 4)

# ══════════════════════════════════════════════════════════════════
# Slide 5 — 지표 수집
# ══════════════════════════════════════════════════════════════════
s = new_slide()
cy = title_block(s, "How The Numbers Are Actually Computed", "지표 수집 — OS 커널에서 직접", sub_line(
    "ResourceCollector(C++, Linux)가 /proc 파일시스템·소켓 옵션을 직접 파싱 — 라이브러리 뒤에 숨지 않고 계산 방식을 직접 결정했다."))
cy += Inches(0.1)
headers = ["지표", "소스", "핵심 포인트"]
col_widths = [Inches(2.0), Inches(2.6), Inches(7.5)]
rows = [
    ["CPU 사용률", "/proc/stat", "누적 jiffies라 두 샘플 사이 델타로 계산"],
    ["메모리 사용량", "/proc/meminfo", "MemFree가 아니라 MemAvailable 사용"],
    ["네트워크 대역폭", "/proc/net/dev", "인터페이스별 누적 바이트 델타, steady_clock으로 실경과시간 측정"],
    ["TCP 연결 품질", "getsockopt(TCP_INFO)", "커널이 이미 추적 중인 RTT/재전송을 그대로 조회"],
    ["디스크 사용량", "statvfs()", "POSIX 표준, 블록 수×블록 크기"],
]
cy = table(s, Inches(0.55), cy, Inches(12.1), headers, rows, col_widths,
           row_h=Inches(0.52), header_h=Inches(0.36), header_size=13, cell_size=14.5)
callout(s, Inches(0.55), Inches(6.95 - 0.85), Inches(12.2), Inches(0.85), sub_line([
    ("왜 MemAvailable인가", TEXT1, True),
    (" — Linux는 남는 메모리를 페이지 캐시로 적극 활용한다. MemFree만 보면 \"거의 다 찼다\"는 착시가 생기는데, "
     "MemAvailable은 즉시 회수 가능한 캐시까지 커널이 계산해 제공하는 값이라 체감 여유 메모리에 더 가깝다.", TEXT2, False),
], size=12.5))
page_num(s, 5)

# ══════════════════════════════════════════════════════════════════
# Slides 6-7 — 설계 결정
# ══════════════════════════════════════════════════════════════════
decisions_p1 = [
    ("01", "Standalone Asio 채택", [
        ("문제", "기존 서버는 Windows IOCP에 강결합 — Linux 등 다른 OS에서 전혀 동작 안 함.", None),
        ("판단", "OS별 비동기 I/O(IOCP/epoll)를 직접 분기하지 않고 이를 추상화하는 Standalone Asio 채택.", None),
        ("해결", "Linux에서 빌드+실행 검증 완료 — 실무 서버 환경에 즉시 적용 가능한 크로스플랫폼 네트워크 계층 확보.", SUCCESS),
    ]),
    ("03", "ARIA-CBC+HMAC → AES-GCM 마이그레이션", [
        ("문제", "수동 조합(Encrypt-then-MAC) 방식이라 순서를 잘못 짜면 패딩 오라클이 재발할 여지.", None),
        ("판단", "AEAD인 GCM은 암호화·인증을 한 번에 처리 — 조합 실수가 구조적으로 불가능. 두 구간 모두 통일.", None),
        ("해결", "ARIA 구현·검증 경험은 코드/문서로 보존, 마이그레이션 후 실제 실행으로 재검증.", SUCCESS),
    ]),
]
decisions_p2 = [
    ("06", "IMetricStore — 컴파일 타임 저장소 선택", [
        ("문제", "SQLite(경량)/TimescaleDB(시계열 특화) 둘 다 지원하고 싶지만, 안 쓰는 백엔드 의존성이 남으면 안 됨.", None),
        ("판단", "런타임 분기 대신 인터페이스 + CMake 옵션으로 컴파일 타임에 하나만 선택.", None),
        ("해결", "ldd로 libpq 미링크 실측 확인 — Postgres 미설치 환경에서도 빌드+실행 가능.", SUCCESS),
    ]),
    ("09", "Windows 포팅 + Smart App Control 대응", [
        ("문제", "빌드는 성공했는데 실행 시 원인 불명으로 멈춤.", None),
        ("판단", "Windows 이벤트 로그로 추적 — Smart App Control이 서명 안 된 vcpkg DLL 로딩을 차단 중이었음. "
                "보안 설정을 낮추는 대신 정적 링크로 전환.", None),
        ("해결", "x64-windows-static로 외부 DLL 0개 완전 정적 빌드 — TLS+AES-GCM 통신·리소스 수집 전체 파이프라인 Windows 재검증.", SUCCESS),
    ]),
]
for page_idx, deck in enumerate([decisions_p1, decisions_p2]):
    s = new_slide()
    cy = title_block(s, None, f"설계 결정 ({page_idx+1}/2)")
    cy += Inches(0.1)
    for num, title_s, rows in deck:
        card_h = decision_card(s, Inches(0.55), cy, Inches(12.2), None, f"{num}. {title_s}", rows)
        cy += card_h + Inches(0.2)
    page_num(s, 6 + page_idx)

# ══════════════════════════════════════════════════════════════════
# Slide 8 — Operational Features
# ══════════════════════════════════════════════════════════════════
s = new_slide()
cy = title_block(s, "Operational Features", '"메트릭 수집기"에서 "APM"으로', sub_line(
    "시스템 리소스만 재는 걸 넘어, 운영 중 실제로 쓰는 4가지 기능을 순서대로 추가."))
cy += Inches(0.1)
features = [
    ("01", "임계치 기반 알림",
     "상태 전이 시에만(OK→Alert→Resolved) 알림 — Zabbix/Nagios 방식. 임계치는 DB 저장+UI 편집.",
     "현재 상태는 매번 DB 재조회로 판단(메모리 캐시 아님) — 프로세스 재시작에도 안 깨짐. 전이 로직은 순수 함수로 분리해 단위 테스트."),
    ("02", "데이터 보존 정책",
     "TimescaleDB는 네이티브 add_retention_policy(), SQLite는 직접 DELETE — 호출부 API는 통일.",
     "Metrics 기본 30일, AlertRecord는 180일 — 백엔드별 구현은 달라도 정책 값은 공유."),
    ("03", "함수/트랜잭션 계측",
     'RAII 스코프 기반 — "우리 APM으로 우리 자신을 계측". 기존 메트릭 전송 파이프라인 재사용.',
     "APM_TRACE_SCOPE 매크로 한 줄로 계측 — Collector.HandleMetricPacket이 첫 적용 사례."),
    ("04", "백분위/집계 통계",
     "SQLite엔 PERCENTILE_CONT 없음 — 시간 창으로 거른 span을 메모리에서 C#으로 직접 계산, 백엔드 무관 통일.",
     "선형 보간 PercentileCalculator(순수 함수) — p50/p95/p99, 시간 창 1h/24h/7d 선택."),
]
fw, fh = Inches(5.95), Inches(2.15)
for i, (num, t, body, detail) in enumerate(features):
    col = i % 2
    row = i // 2
    fx = Inches(0.55) + (fw + Inches(0.3)) * col
    fy = cy + (fh + Inches(0.2)) * row
    rect(s, fx, fy, fw, fh, fill=CARD, line_color=BORDER, line_w=0.75, radius=0.05)
    text(s, fx + Inches(0.22), fy + Inches(0.18), fw - Inches(0.44), Inches(0.32), f"{num}. {t}", 14.5, TEXT1, bold=True, font=FONT_MONO)
    multi_run(s, fx + Inches(0.22), fy + Inches(0.58), fw - Inches(0.44), Inches(0.9),
              sub_line(body, size=11.5), line_spacing=1.3)
    div_y = fy + Inches(1.5)
    divider = s.shapes.add_connector(1, fx + Inches(0.22), div_y, fx + fw - Inches(0.22), div_y)
    divider.line.color.rgb = BORDER2
    divider.line.width = Pt(0.5)
    multi_run(s, fx + Inches(0.22), div_y + Inches(0.14), fw - Inches(0.44), fh - (div_y - fy) - Inches(0.2),
              sub_line([("→ ", ACCENT, True), (detail, TEXT2, False)], size=11), line_spacing=1.28)
page_num(s, 8)

# ══════════════════════════════════════════════════════════════════
# Slide 9 — 300 동시 접속까지
# ══════════════════════════════════════════════════════════════════
s = new_slide()
cy = title_block(s, "Measured, Fixed, Re-measured", "300 동시 접속까지 — 병목을 찾고 고치는 5라운드", sub_line(
    "자체 제작 LoadTester(실제 Agent와 같은 코드 경로, in-process 다중 연결 시뮬레이션)로 매 라운드 재측정."))
cy += Inches(0.1)
headers = ["단계", "100-agent", "300-agent", "비고"]
col_widths = [Inches(2.6), Inches(1.8), Inches(1.8), Inches(5.9)]
rows = [
    [("Round 1(베이스라인)", TEXT2), ("72/100", DANGER), ("71/300", DANGER), ("SQLite fdatasync가 네트워크 스레드 블로킹", TEXT2)],
    [("Round 2(워커 스레드)", TEXT2), ("100/100", SUCCESS), ("229/300", TEXT2), ("저장 비동기화", TEXT2)],
    [("Round 3(WAL)", TEXT2), ("100/100", TEXT2), ("250/300", TEXT2), ("fdatasync↔fcntl 트레이드오프", TEXT2)],
    [("Round 4(로깅 분리)", TEXT2), ("100/100", TEXT2), ("300/300", SUCCESS), ("콘솔 로깅 워커 분리, 전원 성공", TEXT2)],
]
cy = table(s, Inches(0.55), cy, Inches(12.1), headers, rows, col_widths, row_h=Inches(0.4))
cy += Inches(0.25)
callout(s, Inches(0.55), cy, Inches(12.2), Inches(0.65), sub_line([
    ("최종 판단 — 여기서 멈추기로 함", TEXT1, True),
    (" — 남은 병목(futex, 스레드 간 락 경합)은 단일 프로세스·개발 머신 한 대의 현실적 한계로 판단. "
     "실제 운영이라면 Collector 수평 확장이 정공법.", TEXT2, False),
], size=12.5))
cy += Inches(0.85)
bar_chart_two_series(
    s, Inches(0.55), cy, Inches(12.2), "접속 성공률(목표 대비 %) — 라운드별 추이",
    ["Round 1", "Round 2", "Round 3", "Round 4"],
    [72, 100, 100, 100], [24, 76, 83, 100],
    "100-agent", "300-agent", CHART_BLUE, CHART_ORANGE,
)
page_num(s, 9)

# ══════════════════════════════════════════════════════════════════
# Slide 10 — Round 1~2
# ══════════════════════════════════════════════════════════════════
s = new_slide()
cy = title_block(s, None, "Round 1~2 — 72/100 하드 리밋, 그리고 첫 시도의 크래시")
cy += Inches(0.1)
card_h = decision_card(s, Inches(0.55), cy, Inches(12.2), "// Round 1 — 병목 찾기", "72/100에서 하드 리밋", [
    ("결과", "100개 요청 중 72개만 성공, p95 34초.", None),
    ("원인", "strace -c 추적 — fdatasync/pwrite64가 syscall 시간 70%+, SQLite 동기 flush가 네트워크 스레드를 블로킹.", None),
])
cy += card_h + Inches(0.12)
card_h = decision_card(s, Inches(0.55), cy, Inches(12.2), "// Round 2 — 저장 비동기화", "워커 스레드로 위임 — 첫 시도는 크래시", [
    ("문제", "기존 JobQueue 재사용 후 재실측 중 10초 만에 반복 크래시 — 크로스 스레드에서 처음 드러난 락 버그(다음 슬라이드 상세).", None),
    ("해결", "전용 std::mutex 기반 큐로 교체 → 100/100 전원 성공, p95/p99 사실상 0ms. 300-agent도 71→229로 3배 개선.", SUCCESS),
])
cy += card_h + Inches(0.12)
rect(s, Inches(0.55), cy, Inches(12.2), Inches(0.9), fill=CARD, line_color=BORDER, line_w=0.75, radius=0.05)
stat_block(s, Inches(0.95), cy + Inches(0.17), "72 → 100/100", "Round 1 — 100-agent 하드 리밋 해소", w=Inches(5.6))
stat_block(s, Inches(6.75), cy + Inches(0.17), "71 → 229/300", "Round 2 — 300-agent 접속 3배 개선", w=Inches(5.6))
page_num(s, 10)

# ══════════════════════════════════════════════════════════════════
# Slide 11 — Bug spotlight 1
# ══════════════════════════════════════════════════════════════════
s = new_slide()
cy = title_block(s, None, "버그 스포트라이트 — 공유 락 라이브러리의 소유권 버그")
cy += Inches(0.1)
h1 = bug_fix(s, Inches(0.55), cy, Inches(12.2), "bug",
             "문제 — 저장 비동기화 재실측 6단계 전부, 기동 10초 만에 크래시", [
                 ("증상", "GW2_CrossPlatformCore/Thread/JobQueue 재사용 후 재실측 중 모든 단계에서 Collector가 10초 만에 죽었다.", None),
                 ("원인", "WriteUnlock()이 소유자 ID 비트를 안 지움 — 락이 영구 소유된 것처럼 남아 워커 스레드가 10초 스핀 후 크래시.", None),
             ])
cy += h1 + Inches(0.15)
h2 = bug_fix(s, Inches(0.55), cy, Inches(12.2), "fix",
             "해결 — 공유 라이브러리는 안 건드리고, 자체 워커 큐로 의존 자체를 제거", [
                 ("판단", '"검증된 코드"라는 이유로 Lock.cpp는 안 건드리고, std::mutex 기반 전용 WorkerQueue로 JobQueue 의존을 제거.', None),
                 ("검증", "6단계 부하 매트릭스 크래시 없이 완주, 코드도 더 단순해짐.", SUCCESS),
             ])
cy += h2 + Inches(0.15)
code_block(s, Inches(0.55), cy, Inches(12.2), [
    [("// Lock.cpp — WriteUnlock() 내부, 소유자 언락 분기", CODE_CMT)],
    [("if ((threadId & ", None), ("0xFFFF", CODE_TYPE), (") == lockThreadId)", None)],
    [("{", None)],
    [("    _lockFlag.fetch_add(", None), ("1", CODE_TYPE), (");   // 버그: 무관한 하위 비트만 증가", CODE_CMT)],
    [("    return;                   // 소유자 스레드 id 비트는 영원히 안 지워짐", CODE_CMT)],
    [("}", None)],
])
page_num(s, 11)

# ══════════════════════════════════════════════════════════════════
# Slide 12 — Round 3~4
# ══════════════════════════════════════════════════════════════════
s = new_slide()
cy = title_block(s, None, "Round 3~4 — 사라진 줄 알았던 비용, 그리고 다음 층의 병목")
cy += Inches(0.1)
card_h = decision_card(s, Inches(0.55), cy, Inches(12.2), "// Round 3 — \"사라졌다\"는 착시", "strace -f로 다시 보니, 비용은 그대로였다", [
    ("재실측", "strace -c는 워커 스레드를 못 추적 — -f로 재보니 fdatasync 비용이 그대로(전체의 25%), WAL 도입 결정.", None),
    ("결과", "fdatasync -97%, WAL 공유메모리 락(fcntl)은 +437% — fsync 비용이 락 비용으로 형태만 바뀐 트레이드오프.", None),
])
cy += card_h + Inches(0.12)
card_h = decision_card(s, Inches(0.55), cy, Inches(12.2), "// Round 4 — 다음 층의 병목", "이번엔 콘솔 로깅이 1위", [
    ("발견", "300-agent 처리량 증가로 std::cout<<...<<endl(강제 flush)이 write syscall 92%로 새 1위.", None),
    ("해결", "로그를 전용 워커 스레드로 분리 + endl → '\\n' — 300개 전원(250→300) 접속 성공.", SUCCESS),
])
cy += card_h + Inches(0.12)
rect(s, Inches(0.55), cy, Inches(12.2), Inches(0.9), fill=CARD, line_color=BORDER, line_w=0.75, radius=0.05)
stat_block(s, Inches(0.95), cy + Inches(0.17), "-97%", "Round 3 — fdatasync 시간 감소(WAL 적용)", w=Inches(5.6))
stat_block(s, Inches(6.75), cy + Inches(0.17), "300/300", "Round 4 — 300-agent 전원 접속 성공", w=Inches(5.6))
page_num(s, 12)

# ══════════════════════════════════════════════════════════════════
# Slide 13 — Bug spotlight 2
# ══════════════════════════════════════════════════════════════════
s = new_slide()
cy = title_block(s, None, "버그 스포트라이트 — 성능 최적화가 꺼버린 안전장치 (Round 5)")
cy += Inches(0.1)
h1 = bug_fix(s, Inches(0.55), cy, Inches(12.2), "bug",
             "문제 — 로그를 워커로 옮겼더니, 그 워커와 네트워크 스레드가 같은 스트림을 동시에 씀", [
                 ("원인", "sync_with_stdio(false)가 iostream/C stdio 동기화 안전장치를 없앴는데, accept 등 4곳의 cout/cerr가 여전히 네트워크 스레드에 남아 두 스레드가 같은 스트림에 동시 접근.", None),
                 ("발견 계기", "100-agent 벤치마크 진단 로그를 직접 계측하다가 그 로그 자체가 문자 단위로 뒤섞여 나오는 것을 목격.", None),
             ])
cy += h1 + Inches(0.15)
h2 = bug_fix(s, Inches(0.55), cy, Inches(12.2), "fix",
             "수정 — 남은 4곳도 전부 같은 워커로 위임", [
                 ("원칙", '"한 스레드만 스트림을 만진다" 원칙 적용 — sync_with_stdio(false)는 유지, 레이스만 제거.', None),
                 ("검증", "100-agent 스트레스에서 깨진 로그 줄 0건 — 순수 정확성 버그였음을 확인.", SUCCESS),
             ])
cy += h2 + Inches(0.15)
code_block(s, Inches(0.55), cy, Inches(12.2), [
    [("// 수정 전 — 네트워크 스레드에서 cout/cerr를 직접 씀", CODE_CMT)],
    [("if ", CODE_KEY), ("(!ec) { std::cout << ", None), ('"[Collector] connection accepted"', CODE_STR), (" << std::endl; }", None)],
    [("// 수정 후 — 한 워커만 cout/cerr를 건드리도록 통일", CODE_CMT)],
    [("if ", CODE_KEY), ("(!ec) { ", None), ("consoleLogQueue", CODE_TYPE), (".Push([]() { std::cout << ", None),
     ('"[Collector] connection accepted\\n"', CODE_STR), ("; }); }", None)],
])
page_num(s, 13)

# ══════════════════════════════════════════════════════════════════
# Slide 14 — 실측 지표
# ══════════════════════════════════════════════════════════════════
s = new_slide()
cy = title_block(s, "Measured, Not Assumed", "실측 지표", sub_line(
    '"가볍다"는 주장을 수치로 뒷받침 — strace/proc 실측과 테스트 결과.'))
cy += Inches(0.1)
headers = ["항목", "값", "비고"]
col_widths = [Inches(2.6), Inches(2.6), Inches(6.9)]
rows = [
    [("epoll_wait", TEXT2), ("155", TEXT2), ("스핀 없이 블로킹, 사이클당 약 2.6회", TEXT2)],
    [("전체 syscall", TEXT2), ("1,921회 / 19.4ms", TEXT2), ("5분(≈60사이클), 대기가 압도적", TEXT2)],
    [("RSS", SUCCESS), ("13.8MB 고정", SUCCESS), ("15초 간격 20샘플 전부 동일 — 누수 없음", TEXT2)],
    [("CPU 사용률", SUCCESS), ("평균 0.017%", SUCCESS), ("유휴 대기(epoll_wait)가 거의 전부", TEXT2)],
    [("단위 테스트", SUCCESS), ("27 / 27 통과", SUCCESS), ("C++ GoogleTest 9개 + .NET xUnit 18개", TEXT2)],
    [("스케일 테스트", SUCCESS), ("72→300/300 접속", SUCCESS), ("LoadTester 부하 실측 5라운드", TEXT2)],
    [("Windows 빌드/실행", SUCCESS), ("외부 DLL 0개", SUCCESS), ("vcpkg 정적 링크, TLS+AES-GCM 통신 실측 확인", TEXT2)],
]
table(s, Inches(0.55), cy, Inches(12.1), headers, rows, col_widths,
      row_h=Inches(0.54), header_h=Inches(0.36), header_size=13, cell_size=14.5)
page_num(s, 14)

# ══════════════════════════════════════════════════════════════════
# Slides 15-16 — 실제 실행 화면 (2 screens per slide, larger)
# ══════════════════════════════════════════════════════════════════
screens_pages = [
    [("/apm/dashboard", "dashboard.png"), ("/apm/traces", "traces.png")],
    [("/apm/alerts", "alerts.png"), ("터미널 — Collector + LoadTester 300-agent 실행 로그", "terminal.png")],
]
for page_idx, screens in enumerate(screens_pages):
    s = new_slide()
    cy = title_block(s, "Running, Not Just Built", f"실제 실행 화면 ({page_idx + 1}/2)", sub_line(
        "Collector+Console을 함께 띄우고 실 트래픽을 흘려 직접 촬영 — 전부 실제 데이터."))
    cy += Inches(0.1)
    fw, fh = Inches(5.95), Inches(5.0)
    for i, (label, filename) in enumerate(screens):
        fx = Inches(0.55) + (fw + Inches(0.3)) * i
        screenshot_card(s, fx, cy, fw, fh, label, filename)
    page_num(s, 15 + page_idx)

# ══════════════════════════════════════════════════════════════════
# Slide 17 — 기술 스택
# ══════════════════════════════════════════════════════════════════
s = new_slide()
cy = title_block(s, "Tech Stack", "기술 스택")
cy += Inches(0.15)
skill_groups = [
    ("네트워크 코어 (C++)", "C++20, Standalone Asio · OS별 비동기 I/O 추상화(IOCP/epoll) · CMake, GoogleTest, vcpkg · Protobuf 직렬화+리플렉션"),
    ("성능 프로파일링 / 동시성", "strace(-c/-f) syscall 병목 진단 · 부하 테스트 도구 자체 제작(LoadTester) · 워커 큐 설계 · 데이터 레이스 진단/수정"),
    ("보안 설계", "TLS(asio::ssl) · AES-256-GCM(AEAD) · ARIA-256-CBC+HMAC(구현 경험) · 최소 권한(setuid/setgid)"),
    ("웹 대시보드 (.NET)", ".NET 8, ASP.NET Core MVC + Razor · EF Core · SignalR · AssemblyLoadContext 플러그인 · xUnit, Chart.js"),
    ("데이터 / 저장", "SQLite3(WAL) · PostgreSQL/TimescaleDB(선택형) · EF Core 쿼리 번역 한계 진단 + raw SQL 우회"),
]
col_w = Inches(5.95)
col_x = [Inches(0.55), Inches(0.55) + col_w + Inches(0.3)]
col_y = [cy, cy]
col_assign = [0, 1, 0, 1, 0]
for (title_s, body), col in zip(skill_groups, col_assign):
    card_h = tech_card(s, col_x[col], col_y[col], col_w, title_s, body)
    col_y[col] = col_y[col] + card_h + Inches(0.2)
page_num(s, 17)

# ══════════════════════════════════════════════════════════════════
# Slide 18 — Contact
# ══════════════════════════════════════════════════════════════════
s = new_slide()
cy = Inches(1.3)
text(s, Inches(0.55), cy, Inches(11), Inches(0.6), "함께 일하고 싶으시다면", 30, TEXT1, bold=True, font=FONT_MONO)
cy += Inches(0.72)
multi_run(s, Inches(0.55), cy, Inches(11.2), Inches(0.5), sub_line(
    "GitHub 레포지토리에서 소스 코드 전체를 확인하실 수 있습니다 — 궁금한 점은 이메일로 편하게 연락해 주세요.",
    size=15), line_spacing=1.35)
cy += Inches(0.75)

contact_items = [("EMAIL", "shkim4548@gmail.com"), ("GITHUB", "github.com/shkim4548/APM")]
card_w, card_h = Inches(5.95), Inches(1.0)
for i, (label, value) in enumerate(contact_items):
    cx = Inches(0.55) + (card_w + Inches(0.3)) * i
    rect(s, cx, cy, card_w, card_h, fill=CARD, line_color=BORDER, line_w=0.75, radius=0.06)
    text(s, cx + Inches(0.26), cy + Inches(0.18), card_w - Inches(0.52), Inches(0.24), label, 11, TEXT3, font=FONT_MONO)
    text(s, cx + Inches(0.26), cy + Inches(0.46), card_w - Inches(0.52), Inches(0.4), value, 17, ACCENT, font=FONT_MONO)
cy += card_h + Inches(0.55)

divider = s.shapes.add_connector(1, Inches(0.55), cy, Inches(12.75), cy)
divider.line.color.rgb = BORDER2
divider.line.width = Pt(0.75)
cy += Inches(0.35)
final_stats = [("5", "수집 지표"), ("300", "동시 접속 부하 테스트"), ("27", "테스트 통과"), ("10", "버그 발견/수정")]
for i, (val, label) in enumerate(final_stats):
    stat_block(s, Inches(0.55) + Inches(2.95) * i, cy, val, label, w=Inches(2.7))
page_num(s, 18)

out_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "portfolio_apm_light.pptx")
prs.save(out_path)
print("saved:", out_path)
print("slides:", len(prs.slides._sldIdLst))
